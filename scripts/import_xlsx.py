"""Import spreadsheet 'MRFNIP - Cashflow' ke SQLite Monetary.

  python scripts/import_xlsx.py "reference/MRFNIP - Cashflow (2).xlsx" [--reset] [--db data/monetary-v1.db]

Spreadsheet ini memakai bentuk lama (kas & dana darurat terpisah), jadi hasilnya
ditulis ke database v1 lalu dikonversi ke skema aplikasi dengan:

  python scripts/migrate_v2.py --apply --src data/monetary-v1.db

Layout tiap sheet bulanan (mis. NOV-25):
  A-D  : Tanggal, Category, Deskripsi, Amount  (pengeluaran, mulai baris 7)
  F-G  : Deskripsi, Amount                     (pemasukan)
  I-J  : Date/Deskripsi, Amount                (dana darurat, + setor / - tarik)
  L-N  : Category, Sub Category, Amount        (aset; hanya terisi bila ada angka)
  B1   : Saldo sebelumnya  (dipakai untuk bulan pertama)
  J3   : Dana darurat previous (bulan pertama)
"""
import re
import sys
from datetime import datetime, date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import sqlite3  # noqa: E402
from contextlib import contextmanager  # noqa: E402

import openpyxl  # noqa: E402

# Skema v1 sengaja ditulis ulang di sini: aplikasi sudah pakai skema v2, sementara
# spreadsheet sumbernya masih berbentuk lama. Impor -> v1 -> migrate_v2 -> v2.
V1_SCHEMA = """
CREATE TABLE IF NOT EXISTS categories (
    id INTEGER PRIMARY KEY, name TEXT NOT NULL,
    kind TEXT NOT NULL CHECK (kind IN ('expense','income')),
    sort INTEGER NOT NULL DEFAULT 100, active INTEGER NOT NULL DEFAULT 1, UNIQUE(name, kind));
CREATE TABLE IF NOT EXISTS transactions (
    id INTEGER PRIMARY KEY, month_key TEXT NOT NULL,
    kind TEXT NOT NULL CHECK (kind IN ('expense','income')), tx_date TEXT,
    category_id INTEGER REFERENCES categories(id), description TEXT, amount INTEGER NOT NULL,
    status TEXT NOT NULL DEFAULT 'paid', to_fund INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS emergency_fund (
    id INTEGER PRIMARY KEY, month_key TEXT NOT NULL, tx_date TEXT, description TEXT,
    amount INTEGER NOT NULL, linked_tx_id INTEGER);
CREATE TABLE IF NOT EXISTS assets (
    id INTEGER PRIMARY KEY, month_key TEXT NOT NULL, category TEXT NOT NULL, symbol TEXT NOT NULL,
    amount INTEGER NOT NULL, UNIQUE(month_key, category, symbol));
CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT);
"""
DB_OUT = "data/monetary-v1.db"


@contextmanager
def get_db():
    Path(DB_OUT).parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_OUT)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


def init_db():
    with get_db() as db:
        db.executescript(V1_SCHEMA)


def set_setting(db, key, value):
    db.execute("INSERT INTO settings(key,value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
               (key, str(value)))

SHEET_MONTHS = {"JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "MEI": 5, "JUN": 6, "JUL": 7,
                "AUG": 8, "AGU": 8, "SEP": 9, "OKT": 10, "OCT": 10, "NOV": 11, "DES": 12, "DEC": 12}


def sheet_key(title: str):
    m = re.match(r"([A-Z]{3})-(\d{2})", title.strip().upper())
    if not m or m.group(1) not in SHEET_MONTHS:
        return None
    return f"20{m.group(2)}-{SHEET_MONTHS[m.group(1)]:02d}"


def to_int(v):
    if v is None or v == "":
        return None
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def to_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        m = re.search(r"(\d{4}-\d{2}-\d{2})", v)
        if m:
            return m.group(1)
    return None


def split_ef_desc(v):
    """Kolom I bisa berisi tanggal, teks, atau '2025-12-29 [ top up mega ]'."""
    d = to_date(v)
    if isinstance(v, str):
        txt = re.sub(r"\d{4}-\d{2}-\d{2}", "", v).strip(" []")
        return d, (txt or None)
    return d, None


def main(path: str, reset: bool):
    init_db()
    wb = openpyxl.load_workbook(path, data_only=True)
    with get_db() as db:
        if reset:
            for t in ("transactions", "emergency_fund", "assets"):
                db.execute(f"DELETE FROM {t}")
        cat_ids = {(r["name"], r["kind"]): r["id"] for r in db.execute("SELECT id, name, kind FROM categories")}

        def cat(name, kind):
            if not name:
                return None
            name = str(name).strip().upper()
            key = (name, kind)
            if key not in cat_ids:
                cur = db.execute("INSERT INTO categories(name, kind) VALUES (?,?)", (name, kind))
                cat_ids[key] = cur.lastrowid
            return cat_ids[key]

        months = [(sheet_key(ws.title), ws) for ws in wb.worksheets]
        months = sorted([(k, ws) for k, ws in months if k], key=lambda x: x[0])
        first_key = months[0][0]
        first_ws = months[0][1]
        set_setting(db, "first_month", first_key)
        set_setting(db, "opening_balance", str(to_int(first_ws["B1"].value) or 0))
        set_setting(db, "opening_emergency", str(to_int(first_ws["J3"].value) or 0))

        stats = {}
        for mk, ws in months:
            n = [0, 0, 0, 0]
            for r in ws.iter_rows(min_row=7, max_row=500, min_col=1, max_col=4, values_only=True):
                amt = to_int(r[3])
                if not amt:
                    continue
                desc = r[2] if isinstance(r[2], str) else None
                d = to_date(r[0]) or to_date(r[2])
                status = "paid"
                if desc and desc.strip().upper() == "DONE":
                    desc = None
                db.execute(
                    "INSERT INTO transactions(month_key, kind, tx_date, category_id, description, amount, status) VALUES (?,?,?,?,?,?,?)",
                    (mk, "expense", d, cat(r[1], "expense") or cat("OTHER", "expense"), desc.strip() if desc else None, amt, status))
                n[0] += 1
            for r in ws.iter_rows(min_row=7, max_row=100, min_col=6, max_col=7, values_only=True):
                amt = to_int(r[1])
                if not amt:
                    continue
                desc = str(r[0]).strip() if r[0] is not None else None
                k = cat("DANA DARURAT", "income") if desc and "darurat" in desc.lower() else cat("GAJI", "income")
                db.execute(
                    "INSERT INTO transactions(month_key, kind, category_id, description, amount) VALUES (?,?,?,?,?)",
                    (mk, "income", k, desc, amt))
                n[1] += 1
            for r in ws.iter_rows(min_row=7, max_row=99, min_col=9, max_col=10, values_only=True):
                amt = to_int(r[1])
                if not amt:
                    continue
                d, desc = split_ef_desc(r[0])
                db.execute("INSERT INTO emergency_fund(month_key, tx_date, description, amount) VALUES (?,?,?,?)",
                           (mk, d, desc, amt))
                n[2] += 1
            for r in ws.iter_rows(min_row=7, max_row=99, min_col=12, max_col=14, values_only=True):
                amt = to_int(r[2])
                if r[0] is None or amt is None:
                    continue
                db.execute(
                    "INSERT INTO assets(month_key, category, symbol, amount) VALUES (?,?,?,?) "
                    "ON CONFLICT(month_key, category, symbol) DO UPDATE SET amount=excluded.amount",
                    (mk, str(r[0]).strip().upper(), str(r[1] or "").strip().upper(), amt))
                n[3] += 1
            stats[mk] = n

        # PORTOFOLIO = posisi terkini -> snapshot bulan terakhir yang ada transaksinya
        if "PORTOFOLIO" in wb.sheetnames:
            last = max(k for k, v in stats.items() if v[0] or v[1])
            for r in wb["PORTOFOLIO"].iter_rows(min_row=7, max_row=99, min_col=1, max_col=3, values_only=True):
                amt = to_int(r[2])
                if r[0] is None or amt is None:
                    continue
                db.execute(
                    "INSERT INTO assets(month_key, category, symbol, amount) VALUES (?,?,?,?) "
                    "ON CONFLICT(month_key, category, symbol) DO UPDATE SET amount=excluded.amount",
                    (last, str(r[0]).strip().upper(), str(r[1] or "").strip().upper(), amt))

    for mk, (e, i, f, a) in stats.items():
        print(f"{mk}: {e:3d} pengeluaran, {i:2d} pemasukan, {f:2d} dana darurat, {a:2d} aset")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    if "--db" in sys.argv:
        DB_OUT = sys.argv[sys.argv.index("--db") + 1]
    main(sys.argv[1], "--reset" in sys.argv)
    print(f"\nDitulis ke {DB_OUT} (skema lama). Lanjutkan dengan:\n"
          f"  python3 scripts/migrate_v2.py --src {DB_OUT} --apply")
